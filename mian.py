import os
import re
import json
import logging
from datetime import datetime
from telegram import Update
from telegram.ext import Application, CommandHandler, MessageHandler, filters, ContextTypes
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logging.basicConfig(
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    level=logging.INFO
)
logger = logging.getLogger(__name__)

# Sử dụng thư mục /tmp để lưu file (Render có quyền ghi)
DATA_FILE = '/tmp/data_store.json'

def load_data():
    try:
        if os.path.exists(DATA_FILE):
            with open(DATA_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
    except Exception as e:
        logger.error(f"Lỗi khi tải dữ liệu: {e}")
    return {}

def save_data(data):
    try:
        with open(DATA_FILE, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        logger.info(f"Đã lưu dữ liệu: {data}")
    except Exception as e:
        logger.error(f"Lỗi khi lưu dữ liệu: {e}")

# Hàm kiểm tra và tạo file nếu chưa có
def init_data_file():
    if not os.path.exists(DATA_FILE):
        save_data({})
    return load_data()

data_store = init_data_file()

BOT_TOKEN = "8825283140:AAEW53jACQKb1pwGDN5-6ASKKEhWdQf6dvs"

def parse_message(text):
    parts = text.strip().split()
    if len(parts) == 3:
        try:
            float(parts[1].replace(',', ''))
            float(parts[2])
            return {
                "三方": parts[0],
                "金额": parts[1],
                "费率": parts[2],
            }
        except ValueError:
            return None
    return None

def export_to_excel(data_list, date_str):
    filename = f"/tmp/data_{date_str}.xlsx"
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Data {date_str}"
    
    header_font = Font(name='Arial', size=12, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    content_font = Font(name='Arial', size=11)
    content_alignment = Alignment(horizontal='center', vertical='center')
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    headers = ["STT", "三方", "金额", "费率", "时间"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
    
    for row, item in enumerate(data_list, 2):
        ws.cell(row=row, column=1, value=row-1).border = border
        ws.cell(row=row, column=1).alignment = content_alignment
        ws.cell(row=row, column=2, value=item['三方']).border = border
        ws.cell(row=row, column=2).alignment = content_alignment
        ws.cell(row=row, column=2).font = content_font
        ws.cell(row=row, column=3, value=float(item['金额'].replace(',', ''))).border = border
        ws.cell(row=row, column=3).alignment = content_alignment
        ws.cell(row=row, column=3).font = content_font
        ws.cell(row=row, column=3).number_format = '#,##0'
        ws.cell(row=row, column=4, value=float(item['费率'])).border = border
        ws.cell(row=row, column=4).alignment = content_alignment
        ws.cell(row=row, column=4).font = content_font
        ws.cell(row=row, column=4).number_format = '0.00'
        ws.cell(row=row, column=5, value=item['时间']).border = border
        ws.cell(row=row, column=5).alignment = content_alignment
        ws.cell(row=row, column=5).font = content_font
    
    total_row = len(data_list) + 2
    if data_list:
        total_amount = sum(float(item['金额'].replace(',', '')) for item in data_list)
        avg_rate = sum(float(item['费率']) for item in data_list) / len(data_list)
        
        ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=2)
        cell = ws.cell(row=total_row, column=1, value="TỔNG CỘNG")
        cell.font = Font(name='Arial', size=12, bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
        
        ws.cell(row=total_row, column=3, value=total_amount).border = border
        ws.cell(row=total_row, column=3).alignment = content_alignment
        ws.cell(row=total_row, column=3).font = Font(name='Arial', size=11, bold=True)
        ws.cell(row=total_row, column=3).number_format = '#,##0'
        ws.cell(row=total_row, column=4, value=round(avg_rate, 2)).border = border
        ws.cell(row=total_row, column=4).alignment = content_alignment
        ws.cell(row=total_row, column=4).font = Font(name='Arial', size=11, bold=True)
        ws.cell(row=total_row, column=4).number_format = '0.00'
        ws.cell(row=total_row, column=5, value=f"{len(data_list)} records").border = border
        ws.cell(row=total_row, column=5).alignment = content_alignment
        ws.cell(row=total_row, column=5).font = Font(name='Arial', size=11, bold=True)
    
    for col in range(1, len(headers) + 1):
        max_length = 0
        column_letter = get_column_letter(col)
        for row in range(1, total_row + 1):
            cell_value = ws.cell(row=row, column=col).value
            if cell_value:
                length = len(str(cell_value))
                if length > max_length:
                    max_length = length
        adjusted_width = max_length + 4
        ws.column_dimensions[column_letter].width = min(adjusted_width, 30)
    
    wb.save(filename)
    return filename

async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip()
    today_str = datetime.now().strftime("%Y-%m-%d")
    current_time = datetime.now().strftime("%H:%M")
    
    data = parse_message(text)
    if data:
        global data_store
        data_store = load_data()
        if today_str not in data_store:
            data_store[today_str] = []
        data["时间"] = current_time
        data_store[today_str].append(data)
        save_data(data_store)
        
        reply = (f"✅ Đã trích xuất thành công!\n"
                 f"📌 三方: {data['三方']}\n"
                 f"💰 金额: {data['金额']}\n"
                 f"📊 费率: {data['费率']}\n"
                 f"🕐 时间: {data['时间']}\n\n"
                 f"📊 Dùng /export để xuất Excel")
        await update.message.reply_text(reply)

async def today(update: Update, context: ContextTypes.DEFAULT_TYPE):
    today_str = datetime.now().strftime("%Y-%m-%d")
    data_store = load_data()
    if today_str not in data_store or not data_store[today_str]:
        await update.message.reply_text("📭 Chưa có dữ liệu nào trong ngày hôm nay.")
        return
    
    records = data_store[today_str]
    reply = f"📅 Dữ liệu ngày {today_str}:\n\n"
    for idx, item in enumerate(records, 1):
        reply += (f"{idx}. 三方: {item['三方']} | 金额: {item['金额']} | 费率: {item['费率']} | 时间: {item['时间']}\n")
    
    total_amount = sum(float(item['金额'].replace(',', '')) for item in records)
    avg_rate = sum(float(item['费率']) for item in records) / len(records)
    reply += f"\n📊 Tổng số tiền: {total_amount:,.0f}"
    reply += f"\n📊 Trung bình phí: {avg_rate:.2f}"
    reply += f"\n📊 Số lượng giao dịch: {len(records)}"
    await update.message.reply_text(reply)

async def export(update: Update, context: ContextTypes.DEFAULT_TYPE):
    today_str = datetime.now().strftime("%Y-%m-%d")
    data_store = load_data()
    if today_str not in data_store or not data_store[today_str]:
        await update.message.reply_text("📭 Chưa có dữ liệu nào để xuất Excel.")
        return
    
    processing_msg = await update.message.reply_text("⏳ Đang tạo file Excel...")
    try:
        filename = export_to_excel(data_store[today_str], today_str)
        with open(filename, 'rb') as f:
            await update.message.reply_document(
                document=f,
                filename=f"data_{today_str}.xlsx",
                caption=f"📊 Excel export ngày {today_str}\n📈 Tổng số giao dịch: {len(data_store[today_str])}"
            )
        os.remove(filename)
        await processing_msg.delete()
    except Exception as e:
        await processing_msg.edit_text(f"❌ Lỗi khi tạo Excel: {str(e)}")

async def export_all(update: Update, context: ContextTypes.DEFAULT_TYPE):
    data_store = load_data()
    if not data_store:
        await update.message.reply_text("📭 Chưa có dữ liệu nào để xuất.")
        return
    
    processing_msg = await update.message.reply_text("⏳ Đang tạo file Excel tổng hợp...")
    try:
        filename = f"/tmp/data_all_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        wb = openpyxl.Workbook()
        
        for date_str, records in data_store.items():
            if not records:
                continue
            ws = wb.create_sheet(title=date_str[:10])
            headers = ["STT", "三方", "金额", "费率", "时间"]
            for col, header in enumerate(headers, 1):
                ws.cell(row=1, column=col, value=header)
                ws.cell(row=1, column=col).font = Font(bold=True)
            for row, item in enumerate(records, 2):
                ws.cell(row=row, column=1, value=row-1)
                ws.cell(row=row, column=2, value=item['三方'])
                ws.cell(row=row, column=3, value=float(item['金额'].replace(',', '')))
                ws.cell(row=row, column=4, value=float(item['费率']))
                ws.cell(row=row, column=5, value=item['时间'])
            for col in range(1, 6):
                max_length = 0
                for row in range(1, len(records) + 2):
                    cell_value = ws.cell(row=row, column=col).value
                    if cell_value:
                        max_length = max(max_length, len(str(cell_value)))
                ws.column_dimensions[get_column_letter(col)].width = min(max_length + 4, 30)
        
        wb.remove(wb['Sheet'])
        summary_ws = wb.create_sheet(title="Tổng hợp", index=0)
        summary_ws.cell(row=1, column=1, value="Ngày")
        summary_ws.cell(row=1, column=2, value="Số giao dịch")
        summary_ws.cell(row=1, column=3, value="Tổng tiền")
        summary_ws.cell(row=1, column=4, value="Phí trung bình")
        
        row = 2
        for date_str, records in data_store.items():
            if records:
                total = sum(float(item['金额'].replace(',', '')) for item in records)
                avg_rate = sum(float(item['费率']) for item in records) / len(records)
                summary_ws.cell(row=row, column=1, value=date_str)
                summary_ws.cell(row=row, column=2, value=len(records))
                summary_ws.cell(row=row, column=3, value=total)
                summary_ws.cell(row=row, column=4, value=round(avg_rate, 2))
                row += 1
        
        wb.save(filename)
        with open(filename, 'rb') as f:
            await update.message.reply_document(
                document=f,
                filename=f"data_all_{datetime.now().strftime('%Y%m%d')}.xlsx",
                caption=f"📊 Excel tổng hợp tất cả dữ liệu\n📅 Tổng số ngày: {len(data_store)}"
            )
        os.remove(filename)
        await processing_msg.delete()
    except Exception as e:
        await processing_msg.edit_text(f"❌ Lỗi khi tạo Excel: {str(e)}")

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "🤖 **Bot trích xuất dữ liệu DC pha**\n\n"
        "📝 **Cách sử dụng:**\n"
        "Gửi tin nhắn với định dạng:\n"
        "`<Tên> <Số tiền> <Tỷ lệ>`\n\n"
        "**Ví dụ:**\n"
        "`江山 4000 6.9`\n\n"
        "📊 **Lệnh:**\n"
        "/today - Xem dữ liệu hôm nay\n"
        "/export - Xuất Excel hôm nay\n"
        "/export_all - Xuất Excel tất cả dữ liệu\n"
        "/help - Hướng dẫn chi tiết",
        parse_mode='Markdown'
    )

async def help_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "📖 **Hướng dẫn chi tiết**\n\n"
        "**1. Thêm dữ liệu:**\n"
        "Gửi tin nhắn với 3 thông tin:\n"
        "`<Tên> <Số tiền> <Tỷ lệ>`\n"
        "Ví dụ: `江山 4000 6.9`\n\n"
        "**2. Xem dữ liệu:**\n"
        "`/today` - Xem dữ liệu trong ngày\n\n"
        "**3. Xuất Excel:**\n"
        "`/export` - Xuất file Excel của ngày hôm nay\n"
        "`/export_all` - Xuất file Excel tổng hợp tất cả dữ liệu\n\n"
        "**4. Tính năng Excel:**\n"
        "• Định dạng chuyên nghiệp\n"
        "• Tổng hợp số tiền và phí trung bình\n"
        "• Tự động căn chỉnh và tô màu\n"
        "• Sheet tổng hợp khi dùng /export_all",
        parse_mode='Markdown'
    )

def main():
    app = Application.builder().token(BOT_TOKEN).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("help", help_command))
    app.add_handler(CommandHandler("today", today))
    app.add_handler(CommandHandler("export", export))
    app.add_handler(CommandHandler("export_all", export_all))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_message))
    
    port = int(os.environ.get('PORT', 8080))
    webhook_url = f"https://xiafa.onrender.com/{BOT_TOKEN}"
    
    logger.info(f"Starting bot with webhook at {webhook_url}")
    
    app.run_webhook(
        listen='0.0.0.0',
        port=port,
        url_path=BOT_TOKEN,
        webhook_url=webhook_url
    )

if __name__ == "__main__":
    main()
